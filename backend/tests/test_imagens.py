import io
import os
import re
import tempfile

os.environ.setdefault("DATABASE_URL", "sqlite:///" + os.path.join(tempfile.mkdtemp(), "t.db"))

import docx  # noqa: E402
from docx.shared import Cm  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from PIL import Image  # noqa: E402

from app.main import app  # noqa: E402


def _png(w, h, cor="blue"):
    b = io.BytesIO()
    Image.new("RGB", (w, h), cor).save(b, "PNG")
    return b.getvalue()


def test_imagens_upload_e_importacao_docx():
    with TestClient(app) as c:
        tok = c.post("/api/auth/login", json={"login": "admin", "senha": "admin123"}).json()["token"]
        h = {"Authorization": f"Bearer {tok}"}
        # envio direto: imagem grande é reduzida; arquivo que não é imagem é recusado; sem login é recusado
        r = c.post("/api/imagens", files={"arquivo": ("a.png", _png(3200, 1600))}, headers=h)
        assert r.status_code == 200, r.text
        assert r.json()["largura"] == 1600 and r.json()["marcador"].startswith("[imagem:")
        assert c.post("/api/imagens", files={"arquivo": ("x.svg", b"<svg onload=alert(1)>")}, headers=h).status_code == 400
        assert c.post("/api/imagens", files={"arquivo": ("a.png", _png(10, 10))}).status_code == 401
        g = c.get(f"/api/imagens/{r.json()['chave']}")
        assert g.status_code == 200 and g.headers["content-type"] == "image/png"
        assert c.get("/api/imagens/" + "0" * 32).status_code == 404

        # .docx com figura no enunciado e numa alternativa; o "logo" antes da 1ª questão é ignorado
        d = docx.Document()
        d.add_picture(io.BytesIO(_png(100, 40, "gray")), width=Cm(2))
        d.add_paragraph("Prova com figuras")
        d.add_paragraph("1) Observe a figura.")
        d.add_picture(io.BytesIO(_png(400, 200)), width=Cm(8))
        d.add_paragraph("Qual é a cor?")
        d.add_paragraph("a) Azul *")
        d.add_paragraph("b) ")
        d.add_picture(io.BytesIO(_png(50, 50, "red")), width=Cm(1))
        buf = io.BytesIO(); d.save(buf)
        res = c.post("/api/avaliacoes/importar/ler", files={"arquivo": ("p.docx", buf.getvalue())}, headers=h).json()
        q = res["questoes"][0]
        assert re.fullmatch(r"Observe a figura\.\n\[imagem:[0-9a-f]{32}\]\nQual é a cor\?", q["enunciado"])
        assert re.fullmatch(r"\[imagem:[0-9a-f]{32}\]", q["alternativas"][1]["texto"])
        assert q["alternativas"][0]["correta"]
        assert res["titulo_sugerido"] == "Prova com figuras"
        assert any("2 imagem" in a for a in res["avisos"])
