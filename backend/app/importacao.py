"""Leitura de arquivos para importação.

- Alunos: planilhas (.xlsx/.xls), .csv, .txt ou .json em qualquer layout. As colunas RM, nome e
  senha são reconhecidas pelo cabeçalho ou, sem cabeçalho, pelo conteúdo.
- Avaliações: .docx, .pdf ou texto colado, no formato usual de prova (1) enunciado / a) b) c) ...),
  com o gabarito marcado na alternativa, numa linha "Gabarito:" ou numa tabela de gabarito no final.
"""
import csv
import io
import json
import re
import unicodedata

EXTENSOES_TABELA = (".xlsx", ".xlsm", ".xls", ".csv", ".txt", ".tsv", ".json")
EXTENSOES_PROVA = (".docx", ".pdf", ".txt", ".md")


class ErroImportacao(ValueError):
    pass


def _norm(s) -> str:
    s = unicodedata.normalize("NFKD", str(s or ""))
    s = "".join(c for c in s if not unicodedata.combining(c)).lower()
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def _celula(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():  # Excel guarda 12345 como 12345.0
        v = int(v)
    return str(v).strip()


def _decodificar(dados: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return dados.decode(enc)
        except UnicodeDecodeError:
            continue
    return dados.decode("utf-8", errors="replace")


# ======================================================================
# ALUNOS
# ======================================================================
def ler_tabela(nome_arquivo: str, dados: bytes) -> list[list[str]]:
    """Devolve as linhas do arquivo como lista de listas de texto."""
    nome = (nome_arquivo or "").lower()
    if nome.endswith((".xlsx", ".xlsm")):
        import openpyxl
        try:
            wb = openpyxl.load_workbook(io.BytesIO(dados), read_only=True, data_only=True)
        except Exception:
            raise ErroImportacao("Não foi possível abrir a planilha. Salve como .xlsx ou .csv e tente de novo.")
        for ws in wb.worksheets:  # primeira aba que tiver dados
            linhas = [[_celula(v) for v in row] for row in ws.iter_rows(values_only=True)]
            linhas = [l for l in linhas if any(l)]
            if linhas:
                return linhas
        return []
    if nome.endswith(".xls"):
        import xlrd
        try:
            wb = xlrd.open_workbook(file_contents=dados)
        except Exception:
            raise ErroImportacao("Não foi possível abrir a planilha .xls. Salve como .xlsx ou .csv.")
        for ws in wb.sheets():
            linhas = [[_celula(ws.cell_value(r, c)) for c in range(ws.ncols)] for r in range(ws.nrows)]
            linhas = [l for l in linhas if any(l)]
            if linhas:
                return linhas
        return []
    texto = _decodificar(dados)
    if nome.endswith(".json") or texto.lstrip()[:1] in ("[", "{"):
        return _tabela_json(texto)
    return _tabela_texto(texto)


def _tabela_json(texto: str) -> list[list[str]]:
    try:
        obj = json.loads(texto)
    except json.JSONDecodeError as e:
        raise ErroImportacao(f"JSON inválido (linha {e.lineno}): {e.msg}")
    if isinstance(obj, dict):  # {"alunos": [...]} ou {"turma": ..., "lista": [...]}
        listas = [v for v in obj.values() if isinstance(v, list)]
        if not listas:
            raise ErroImportacao("O JSON precisa conter uma lista de alunos")
        obj = max(listas, key=len)
    if not isinstance(obj, list) or not obj:
        raise ErroImportacao("O JSON precisa conter uma lista de alunos")
    if all(isinstance(i, dict) for i in obj):
        chaves: list[str] = []
        for item in obj:
            for k in item:
                if k not in chaves:
                    chaves.append(k)
        return [chaves] + [[_celula(item.get(k)) for k in chaves] for item in obj]
    return [[_celula(v) for v in (i if isinstance(i, list) else [i])] for i in obj]


def _tabela_texto(texto: str) -> list[list[str]]:
    linhas = [l for l in texto.splitlines() if l.strip()]
    if not linhas:
        return []
    amostra = "\n".join(linhas[:20])
    contagem = {d: sum(l.count(d) for l in linhas[:20]) for d in (";", "\t", ",", "|")}
    delim = max(contagem, key=contagem.get)
    if contagem[delim] == 0:
        # sem separador: tenta "RM  Nome Completo" separado por 2+ espaços
        return [re.split(r"\s{2,}", l.strip()) for l in linhas]
    try:
        delim = csv.Sniffer().sniff(amostra, delimiters=";\t,|").delimiter
    except csv.Error:
        pass
    return [[c.strip() for c in row] for row in csv.reader(linhas, delimiter=delim) if any(c.strip() for c in row)]


SINONIMOS = {
    "rm": {"rm", "r m", "ra", "r a", "matricula", "n matricula", "numero matricula", "registro", "registro aluno",
           "codigo", "cod", "codigo aluno", "login", "usuario", "user", "id", "id aluno", "numero", "n", "no",
           "cgm", "rga", "registro academico"},
    "nome": {"nome", "aluno", "nome aluno", "nome do aluno", "nome completo", "estudante", "nome estudante",
             "discente", "name", "nome social"},
    "senha": {"senha", "password", "pass", "pwd", "senha inicial", "senha acesso", "palavra passe"},
}


def _campo_do_cabecalho(texto: str) -> str | None:
    n = _norm(texto)
    for campo, nomes in SINONIMOS.items():
        if n in nomes:
            return campo
    for campo, nomes in SINONIMOS.items():  # "RM do aluno", "Nome completo do estudante"
        if any(n.startswith(x + " ") for x in nomes if len(x) > 1):
            return campo
    return None


def detectar_colunas(linhas: list[list[str]]) -> dict:
    """Descobre quais colunas são RM, nome e senha. Devolve índices (ou None) e se há cabeçalho."""
    if not linhas:
        return {"tem_cabecalho": False, "rm": None, "nome": None, "senha": None}
    largura = max(len(l) for l in linhas)
    cab = linhas[0]
    mapa = {"rm": None, "nome": None, "senha": None}
    fracos = {"n", "no", "numero", "id", "cod", "codigo"}  # "Nº" da chamada perde para "RA"/"RM"
    prioridade = {}
    for i, c in enumerate(cab):
        campo = _campo_do_cabecalho(c)
        if not campo:
            continue
        p = 0 if _norm(c) in fracos else 1
        if mapa[campo] is None or p > prioridade[campo]:
            mapa[campo], prioridade[campo] = i, p
    tem_cab = sum(v is not None for v in mapa.values()) >= 1 and not any(
        re.fullmatch(r"\d{3,}", c or "") for c in cab)
    if tem_cab and mapa["rm"] is not None and mapa["nome"] is not None:
        return {"tem_cabecalho": True, **mapa}

    # heurística pelo conteúdo
    dados = linhas[1:] if tem_cab else linhas
    amostra = dados[:200]

    def taxa(i, teste):
        vals = [l[i] for l in amostra if i < len(l) and l[i]]
        return (sum(1 for v in vals if teste(v)) / len(vals)) if vals else 0

    def parece_nome(v):
        return bool(re.search(r"[A-Za-zÀ-ÿ]{2,}\s+[A-Za-zÀ-ÿ]{2,}", v)) and not re.search(r"\d{3,}", v)

    def parece_rm(v):
        return bool(re.fullmatch(r"[A-Za-z]{0,3}[\d.\-/]{2,}[A-Za-z0-9]{0,3}", v.replace(" ", "")))

    livres = [i for i in range(largura) if i not in mapa.values()]
    if mapa["nome"] is None and livres:
        melhor = max(livres, key=lambda i: taxa(i, parece_nome))
        if taxa(melhor, parece_nome) >= 0.5:
            mapa["nome"] = melhor
            livres.remove(melhor)
    if mapa["rm"] is None and livres:
        melhor = max(livres, key=lambda i: taxa(i, parece_rm))
        if taxa(melhor, parece_rm) >= 0.5:
            mapa["rm"] = melhor
            livres.remove(melhor)
    if mapa["senha"] is None and livres and not tem_cab:
        # uma coluna sobrando e preenchida, sem cabeçalho: provavelmente a senha
        cheias = [i for i in livres if taxa(i, lambda v: True) > 0 and
                  sum(1 for l in amostra if i < len(l) and l[i]) >= 0.8 * len(amostra)]
        if len(cheias) == 1:
            mapa["senha"] = cheias[0]
    return {"tem_cabecalho": tem_cab, **mapa}


def normalizar_rm(v: str) -> str:
    v = (v or "").strip()
    if re.fullmatch(r"\d+\.0", v):
        v = v[:-2]
    return v


# ======================================================================
# AVALIAÇÕES
# ======================================================================
MARCA_IMG_ARQUIVO = re.compile(r"\[imagem-arquivo:(\d+)\]")


def extrair_linhas(nome_arquivo: str, dados: bytes, imagens: list | None = None) -> list[tuple[str, bool]]:
    """Linhas de texto do documento. O booleano indica se a linha inteira está em negrito (docx).

    Se `imagens` for uma lista, as figuras do .docx são acrescentadas a ela (bytes) e, no texto,
    aparece o marcador [imagem-arquivo:N] na posição da figura.
    """
    nome = (nome_arquivo or "").lower()
    if nome.endswith(".docx"):
        import docx
        try:
            d = docx.Document(io.BytesIO(dados))
        except Exception:
            raise ErroImportacao("Não foi possível abrir o .docx. Se for .doc antigo, salve como .docx no Word.")
        out = []

        def figuras(p):
            if imagens is None:
                return []
            marcas = []
            try:
                ids = p._p.xpath(".//a:blip/@r:embed")
            except Exception:
                ids = []
            for rid in ids:
                try:
                    parte = d.part.related_parts[rid]
                    blob = parte.blob
                except Exception:
                    continue
                if parte.content_type in ("image/x-emf", "image/x-wmf", "image/emf", "image/wmf"):
                    continue  # formatos do Windows que o navegador não exibe
                imagens.append(blob)
                marcas.append(f"[imagem-arquivo:{len(imagens) - 1}]")
            return marcas

        def par(p):
            runs = [r for r in p.runs if r.text.strip()]
            negrito = bool(runs) and all(r.bold or (r.style and r.style.font.bold) for r in runs)
            if p.text.strip():
                for pedaco in p.text.split("\n"):
                    out.append((pedaco, negrito))
            for marca in figuras(p):
                out.append((marca, False))

        corpo = d.element.body
        for bloco in corpo.iterchildren():
            tag = bloco.tag.rsplit("}", 1)[-1]
            if tag == "p":
                par(docx.text.paragraph.Paragraph(bloco, d))
            elif tag == "tbl":  # tabelas (ex.: gabarito em tabela)
                t = docx.table.Table(bloco, d)
                for row in t.rows:
                    celulas = []
                    for c in row.cells:
                        if c.text.strip() and (not celulas or celulas[-1] != c.text.strip()):
                            celulas.append(c.text.strip())
                    out.append(("  ".join(celulas), False))
        return out
    if nome.endswith(".pdf"):
        from pypdf import PdfReader
        try:
            leitor = PdfReader(io.BytesIO(dados))
            texto = "\n".join((pg.extract_text() or "") for pg in leitor.pages)
        except Exception:
            raise ErroImportacao("Não foi possível ler o PDF.")
        if len(texto.strip()) < 20:
            raise ErroImportacao("O PDF não tem texto selecionável (provavelmente é uma imagem escaneada). "
                                 "Use o arquivo original do Word, ou copie e cole o texto.")
        return [(l, False) for l in texto.splitlines()]
    if nome.endswith(".doc"):
        raise ErroImportacao("Arquivos .doc antigos não são suportados. Abra no Word e salve como .docx.")
    return [(l, False) for l in _decodificar(dados).splitlines()]


RE_QUESTAO = re.compile(r"^\s*(?:quest[aã]o|q\.?)?\s*(\d{1,3})\s*(?:[\)\.\-–—:º°]|\s*-)\s*(.*)$", re.I)
RE_QUESTAO_TITULO = re.compile(r"^\s*quest[aã]o\s*(\d{1,3})\s*[:\-–—.]?\s*(.*)$", re.I)
RE_ALT = re.compile(r"^\s*[\(\[]?([a-jA-J])\s*[\)\]\.\-–—]\s*(.*)$")
RE_GAB_LINHA = re.compile(r"^\s*(?:gabarito|resposta(?:\s+(?:correta|esperada|modelo|certa))?|alternativa\s+correta|correta|"
                          r"expectativa\s+de\s+resposta)\s*[:\-–—=]\s*(.+)$", re.I)
RE_CHAVES = re.compile(r"^\s*palavras?[\s\-]*chaves?\s*[:\-–—]\s*(.+)$", re.I)
RE_PESO = re.compile(r"^\s*(?:peso|valor|pontos?)\s*[:\-–—=]\s*([\d]+(?:[.,]\d+)?)", re.I)
RE_EXPL = re.compile(r"^\s*(?:explica[cç][aã]o|coment[aá]rio|justificativa|resolu[cç][aã]o)\s*[:\-–—]\s*(.+)$", re.I)
RE_SECAO_GAB = re.compile(r"^\s*(?:gabarito|respostas|folha\s+de\s+respostas)\s*(?:oficial)?\s*:?\s*$", re.I)
RE_PAR_GAB = re.compile(r"(\d{1,3})\s*[\-\)\.:–—=]?\s*\(?([A-Ja-j])\)?(?![A-Za-zÀ-ÿ])")
RE_SECAO = re.compile(r"^\s*(?:(?:parte|se[cç][aã]o|bloco|etapa)\s+(?:[ivxlc]+|\d+|[a-d])\b|boa\s+(?:prova|sorte)\b|fim\s+da\s+(?:prova|avalia))", re.I)
MARCA_CORRETA = re.compile(r"(\s*[\(\[]\s*[xX✓✔]\s*[\)\]]\s*|\s*\*+\s*|\s*[✓✔]\s*|\s*\((?:correta|certa)\)\s*)")


def _tirar_marca(texto: str) -> tuple[str, bool]:
    t = texto.strip()
    marcado = False
    for _ in range(2):
        m_ini = re.match(r"^(?:[\(\[]\s*[xX✓✔]\s*[\)\]]|\*+|[✓✔])\s*", t)
        m_fim = re.search(r"\s*(?:[\(\[]\s*[xX✓✔]\s*[\)\]]|\*+|[✓✔]|\((?:correta|certa)\))$", t, re.I)
        if m_ini:
            t, marcado = t[m_ini.end():], True
        if m_fim:
            t, marcado = t[:m_fim.start()], True
    return t.strip(), marcado


def interpretar_prova(linhas: list[tuple[str, bool]]) -> dict:
    questoes: list[dict] = []
    cabecalho: list[str] = []
    atual = None
    campo = None           # onde vão as linhas de continuação: enunciado | alt | gabarito | explicacao
    gabarito_final: dict[int, str] = {}
    na_secao_gab = False

    def nova(numero, texto):
        return {"numero": numero, "enunciado": texto, "alternativas": [], "gabarito_texto": "",
                "palavras_chave": [], "peso": None, "explicacao": "", "letra_correta": None,
                "tipo_forcado": None}

    for bruto, negrito in linhas:
        linha = bruto.replace(" ", " ").rstrip()
        if not linha.strip():
            continue
        if re.fullmatch(r"\s*[-_=*.]{3,}\s*", linha):
            continue

        # seção de gabarito no final da prova (uma grade "GABARITO" em branco antes das questões é ignorada)
        if RE_SECAO_GAB.match(linha):
            if atual is not None:
                na_secao_gab = True
            continue
        # títulos de seção ("PARTE II: Questões Dissertativas", "Boa prova!") não pertencem a nenhuma questão
        if atual is not None and RE_SECAO.match(linha):
            campo = "ignorar"
            continue
        pares = RE_PAR_GAB.findall(linha)
        if na_secao_gab or (len(pares) >= 3 and re.match(r"^\s*(gabarito|respostas)?\s*:?\s*\d", linha, re.I)):
            if pares:
                for n, l in pares:
                    gabarito_final[int(n)] = l.upper()
                continue
            if na_secao_gab and not RE_QUESTAO.match(linha):
                continue
        m_gab = RE_GAB_LINHA.match(linha)
        if m_gab and (len(RE_PAR_GAB.findall(m_gab.group(1))) >= 2 or
                      re.match(r"^\s*\d{1,3}\s*[\-\)\.:–—=]\s*\(?[A-Ja-j]\)?\s*(?:$|[,;\s])", m_gab.group(1))):
            for n, l in RE_PAR_GAB.findall(m_gab.group(1)):
                gabarito_final[int(n)] = l.upper()
            continue

        m_q = RE_QUESTAO_TITULO.match(linha) or RE_QUESTAO.match(linha)
        if m_q and not (atual and campo == "alt" and RE_ALT.match(linha)):
            numero, texto = int(m_q.group(1)), m_q.group(2).strip()
            esperado = (atual["numero"] + 1) if atual else None
            # evita confundir "1990. Em..." ou números soltos dentro do texto com nova questão
            if esperado is None or numero == esperado or numero == 1 or RE_QUESTAO_TITULO.match(linha):
                if atual:
                    questoes.append(atual)
                atual = nova(numero, texto)
                campo = "enunciado"
                continue
        if atual is None:
            cabecalho.append(linha.strip())
            continue
        if campo == "ignorar" and not (RE_GAB_LINHA.match(linha) or RE_CHAVES.match(linha)):
            continue

        if (m := RE_CHAVES.match(linha)):
            atual["palavras_chave"] += [p.strip() for p in re.split(r"[;,]", m.group(1)) if p.strip()]
            atual["tipo_forcado"] = "dissertativa"
            campo = None
            continue
        if (m := RE_PESO.match(linha)):
            atual["peso"] = float(m.group(1).replace(",", "."))
            continue
        if (m := RE_EXPL.match(linha)):
            atual["explicacao"] = m.group(1).strip()
            campo = "explicacao"
            continue
        if (m := RE_GAB_LINHA.match(linha)):
            valor = m.group(1).strip()
            letra = re.fullmatch(r"\(?([A-Ja-j])\)?[\s.]*(?:\).*)?", valor)
            if atual["alternativas"] and letra:
                atual["letra_correta"] = letra.group(1).upper()
            else:
                atual["gabarito_texto"] = valor
                atual["tipo_forcado"] = atual["tipo_forcado"] or ("dissertativa" if not atual["alternativas"] else None)
                campo = "gabarito"
            continue
        m_a = RE_ALT.match(linha)
        if m_a and (atual["alternativas"] or m_a.group(1).lower() == "a") and campo != "gabarito":
            letra = m_a.group(1).upper()
            esperado = chr(ord("A") + len(atual["alternativas"]))
            if letra == esperado:
                texto, marcada = _tirar_marca(m_a.group(2))
                atual["alternativas"].append({"texto": texto, "correta": marcada or negrito})
                campo = "alt"
                continue
        # continuação
        if campo == "alt" and atual["alternativas"]:
            atual["alternativas"][-1]["texto"] += " " + linha.strip()
        elif campo == "gabarito":
            atual["gabarito_texto"] += "\n" + linha.rstrip()  # mantém quebras e recuos (ex.: código)
        elif campo == "explicacao":
            atual["explicacao"] += " " + linha.strip()
        else:
            atual["enunciado"] += ("\n" if atual["enunciado"] else "") + linha.rstrip()
    if atual:
        questoes.append(atual)

    # monta o resultado final
    saida, avisos = [], []
    for q in questoes:
        enun = q["enunciado"].strip()
        if re.search(r"\((?:dissertativa|discursiva|aberta)\)", enun, re.I):
            q["tipo_forcado"] = "dissertativa"
            enun = re.sub(r"\s*\((?:dissertativa|discursiva|aberta)\)\s*", " ", enun, flags=re.I).strip()
        enun = re.sub(r"\s*\((?:(?:valor|peso)\s*:?\s*(\d+(?:[.,]\d+)?)\s*(?:pontos?|pts?)?|(\d+(?:[.,]\d+)?)\s*(?:pontos?|pts?))\)",
                      lambda m: (q.__setitem__("peso", q["peso"] or float((m.group(1) or m.group(2)).replace(",", "."))) or ""),
                      enun, flags=re.I).strip()
        alts = q["alternativas"]
        tipo = "objetiva" if len(alts) >= 2 and q["tipo_forcado"] != "dissertativa" else "dissertativa"
        item = {"numero": q["numero"], "tipo": tipo, "enunciado": enun, "peso": q["peso"],
                "explicacao": q["explicacao"].strip() or None, "avisos": []}
        if tipo == "objetiva":
            letra = q["letra_correta"] or gabarito_final.get(q["numero"])
            marcadas = [i for i, a in enumerate(alts) if a["correta"]]
            if letra and ord(letra) - 65 < len(alts):
                correta = ord(letra) - 65
            elif len(marcadas) == 1:
                correta = marcadas[0]
            else:
                correta = None
                item["avisos"].append("Gabarito não encontrado — marque a alternativa correta"
                                      if not marcadas else "Mais de uma alternativa marcada — escolha a correta")
            item["alternativas"] = [{"texto": a["texto"], "correta": i == correta} for i, a in enumerate(alts)]
        else:
            item["gabarito_texto"] = q["gabarito_texto"].strip() or None
            item["palavras_chave"] = q["palavras_chave"]
            item["alternativas"] = []
            if not item["gabarito_texto"] and not item["palavras_chave"]:
                item["avisos"].append("Sem resposta esperada/palavras-chave — a correção será manual")
        if not enun:
            item["avisos"].append("Enunciado vazio")
        for a in item["avisos"]:
            avisos.append(f"Questão {q['numero']}: {a}")
        saida.append(item)

    if not saida:
        raise ErroImportacao("Nenhuma questão encontrada. As questões devem começar com número, ex.: \"1) Qual...\" "
                             "ou \"Questão 1\", e as alternativas com letra, ex.: \"a) ...\". Baixe o modelo para ver o formato.")
    cabecalho = [c for c in cabecalho if not MARCA_IMG_ARQUIVO.fullmatch(c.strip())]
    candidatos = [c for c in cabecalho if len(c) > 3 and not re.match(r"^(nome|aluno|turma|data|escola|professor|docente|ra\b)", c, re.I)]
    titulo = next((c for c in candidatos if re.search(r"avalia|prova|simulado|exame|teste|\bp\d\b", c, re.I)),
                  candidatos[0] if candidatos else None)
    return {"titulo_sugerido": titulo, "cabecalho": cabecalho[:10], "questoes": saida, "avisos": avisos}
